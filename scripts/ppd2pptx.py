# -*- coding: utf-8 -*-
"""
pptd/.page (960x540 声明式幻灯片 DSL)  ->  PPTX 转换器

坐标映射：1 px = 1 pt = 12700 EMU，画布 960x540 px -> 960pt x 540pt (13.333in x 7.5in)
"""
import os
import re
import html
import math
import copy
import sys
import yaml

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from lxml import etree

PX = 12700  # EMU per px(=pt)

# ---------------------------------------------------------------- helpers

def E(v):
    return Emu(int(round(v * PX)))


def OxmlElement(tag, nsdecls=None):
    from pptx.oxml import parse_xml as _p
    from pptx.oxml.ns import nsdecls as _nd
    if nsdecls is None:
        nsdecls = _nd('a')
    return _p('<%s %s/>' % (tag, nsdecls))


RPR_ORDER = ['a:ln', 'a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill',
             'a:grpFill', 'a:effectLst', 'a:effectDag', 'a:highlight', 'a:uLnTx', 'a:uLn',
             'a:uFillTx', 'a:uFill', 'a:latin', 'a:ea', 'a:cs', 'a:sym', 'a:hlinkClick',
             'a:hlinkMouseOver', 'a:rtl', 'a:extLst']


def _insert_ordered(parent, el, order):
    tag = etree.QName(el).localname
    tagname = 'a:' + tag
    idx = order.index(tagname) if tagname in order else len(order)
    for child in parent:
        cname = 'a:' + etree.QName(child).localname
        cidx = order.index(cname) if cname in order else len(order)
        if cidx > idx:
            child.addprevious(el)
            return
    parent.append(el)


def set_run_font(run, latin=None, ea=None, cs=None):
    rPr = run.font._rPr
    for tagname, val in (('a:latin', latin), ('a:ea', ea), ('a:cs', cs)):
        if val is None:
            continue
        old = rPr.find(qn(tagname))
        if old is not None:
            rPr.remove(old)
        el = OxmlElement(tagname)
        el.set('typeface', val)
        _insert_ordered(rPr, el, RPR_ORDER)


def set_run_spacing(run, px_val):
    rPr = run.font._rPr
    rPr.set('spc', str(int(round(px_val * 100))))


def set_run_baseline(run, permille):
    rPr = run.font._rPr
    rPr.set('baseline', str(int(permille)))


def no_line(shape):
    shape.line.fill.background()


# ---------------------------------------------------------------- theme

class Theme(object):
    def __init__(self, d):
        self.raw = d
        self.colors = d.get('theme', {}).get('colors', {})
        self.textStyles = d.get('theme', {}).get('textStyles', {})
        self.tableStyles = d.get('theme', {}).get('tableStyles', {})

    def color(self, v):
        if v is None:
            return None
        if isinstance(v, str) and v.startswith('$'):
            c = self.colors.get(v[1:])
            if c is None:
                raise KeyError('unknown color %s' % v)
            return c
        return v

    def text_style(self, name):
        key = name[1:] if isinstance(name, str) and name.startswith('$') else name
        st = self.textStyles.get(key)
        return copy.deepcopy(st) if st else {}


EA_MAP = {
    '思源宋体': '宋体',
    'Source Han Serif': '宋体',
    'Noto Serif CJK SC': '宋体',
}


def norm_ea(name):
    if not name:
        return '宋体'
    return EA_MAP.get(name, name)


def split_family(fam):
    """返回 (latin, ea)"""
    if fam is None:
        return ('Times New Roman', '宋体')
    if isinstance(fam, dict):
        return (fam.get('latin', 'Times New Roman'), norm_ea(fam.get('ea')))
    if isinstance(fam, str):
        return (fam, '宋体')
    return ('Times New Roman', '宋体')


ALIGN_H = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT,
           'justify': PP_ALIGN.JUSTIFY}
ALIGN_V = {'top': MSO_ANCHOR.TOP, 'middle': MSO_ANCHOR.MIDDLE, 'bottom': MSO_ANCHOR.BOTTOM}


# ---------------------------------------------------------------- LaTeX

SYMBOLS = {
    'ell': 'ℓ', 'times': '×', 'cdot': '·', 'sum': 'Σ', 'prod': '∏', 'int': '∫',
    'alpha': 'α', 'beta': 'β', 'gamma': 'γ', 'delta': 'δ', 'Delta': 'Δ', 'epsilon': 'ε',
    'varepsilon': 'ε', 'zeta': 'ζ', 'eta': 'η', 'theta': 'θ', 'Theta': 'Θ', 'iota': 'ι',
    'kappa': 'κ', 'lambda': 'λ', 'Lambda': 'Λ', 'mu': 'μ', 'nu': 'ν', 'xi': 'ξ', 'pi': 'π',
    'Pi': 'Π', 'rho': 'ρ', 'sigma': 'σ', 'Sigma': 'Σ', 'tau': 'τ', 'upsilon': 'υ',
    'phi': 'φ', 'Phi': 'Φ', 'chi': 'χ', 'psi': 'ψ', 'omega': 'ω', 'Omega': 'Ω',
    'le': '≤', 'ge': '≥', 'neq': '≠', 'approx': '≈', 'pm': '±', 'infty': '∞',
    'to': '→', 'rightarrow': '→', 'leftarrow': '←', 'Rightarrow': '⇒', 'in': '∈',
    'partial': '∂', 'nabla': '∇', 'prime': '′', 'ldots': '…', 'dots': '…',
}

SPACES = {'\\,': '', '\\;': '', '\\:': '', '\\!': '', '\\ ': ' ',
          '\\qquad': '\u2003', '\\quad': '\u2002'}


def _read_group(s, i):
    """s[i] == '{' -> 返回 (内容, 新下标)"""
    depth = 0
    for j in range(i, len(s)):
        if s[j] == '{':
            depth += 1
        elif s[j] == '}':
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
    return s[i + 1:], len(s)


def _runs_text(runs):
    return ''.join(r['text'] for r in runs)


def latex_to_runs(s, italic=True, tight=False):
    """把简单 LaTeX 片段转成 run 列表：{text, italic, bold, sub, sup}"""
    out = []

    def push(text, **kw):
        if text == '':
            return
        d = {'text': text}
        d.update(kw)
        out.append(d)

    def last_ch():
        for r in reversed(out):
            if r.get('text'):
                return r['text'][-1]
        return None

    def push_op(op):
        prev = last_ch()
        if tight:
            push(op)
            return
        if prev is None or prev in '([{ =+−×·／,<≥≤':
            push(op)
        else:
            push(' ' + op + ' ')

    def push_runs(runs, sub=False, sup=False):
        for r in runs:
            d = dict(r)
            if sub:
                d['sub'] = True
            if sup:
                d['sup'] = True
            out.append(d)

    i = 0
    n = len(s)
    while i < n:
        c = s[i]
        if c == '\\':
            if s[i:i + 2] in SPACES:
                push(SPACES[s[i:i + 2]])
                i += 2
                continue
            m = re.match(r'\\([A-Za-z]+)', s[i:])
            if not m:
                i += 1
                continue
            cmd = m.group(1)
            i += 1 + len(cmd)
            if cmd in ('text', 'mathrm', 'mathbf', 'mathit', 'operatorname'):
                if i < n and s[i] == '{':
                    inner, i = _read_group(s, i)
                else:
                    inner = s[i]
                    i += 1
                inner = re.sub(r'\\([A-Za-z]+)', lambda mm: SYMBOLS.get(mm.group(1), mm.group(1)), inner)
                push(inner, bold=(cmd == 'mathbf'), italic=(cmd == 'mathit') and italic)
                continue
            if cmd in ('frac', 'dfrac', 'tfrac'):
                num, i = _read_group(s, i)
                den, i = _read_group(s, i)
                nr = latex_to_runs(num, italic)
                dr = latex_to_runs(den, italic)
                if _needs_paren(num):
                    push('（')
                push_runs(nr)
                if _needs_paren(num):
                    push('）')
                push('／')
                if _needs_paren(den):
                    push('（')
                push_runs(dr)
                if _needs_paren(den):
                    push('）')
                continue
            if cmd in ('left', 'right'):
                if i < n:
                    d = s[i]
                    i += 1
                    if d == '[':
                        push('[')
                    elif d == ']':
                        push(']')
                    elif d == '(':
                        push('(')
                    elif d == ')':
                        push(')')
                    elif d == '|':
                        push('|')
                    elif d == '.':
                        pass
                    else:
                        push(d)
                continue
            if cmd == 'bar':
                if i < n and s[i] == '{':
                    g, i = _read_group(s, i)
                else:
                    g = s[i]
                    i += 1
                push_runs(latex_to_runs(g, italic))
                # 组合长音符
                if out:
                    out[-1]['text'] = out[-1]['text'] + '\u0304'
                continue
            if cmd in ('hat', 'tilde', 'vec'):
                if i < n and s[i] == '{':
                    g, i = _read_group(s, i)
                else:
                    g = s[i]
                    i += 1
                mark = {'hat': '\u0302', 'tilde': '\u0303', 'vec': '\u20D7'}[cmd]
                push_runs(latex_to_runs(g, italic))
                if out:
                    out[-1]['text'] = out[-1]['text'] + mark
                continue
            if cmd in SYMBOLS:
                if cmd in ('times', 'cdot', 'le', 'ge', 'neq', 'approx', 'pm'):
                    push_op(SYMBOLS[cmd])
                else:
                    push(SYMBOLS[cmd], italic=False)
                continue
            push(cmd)
            continue
        if c in '_^':
            sub = (c == '_')
            i += 1
            if i < n and s[i] == '{':
                g, i = _read_group(s, i)
            elif i < n:
                g = s[i]
                i += 1
            else:
                g = ''
            push_runs(latex_to_runs(g, italic, tight=True), sub=sub, sup=not sub)
            continue
        if c == ' ':
            push(' ')
            i += 1
            continue
        if c == '-':
            push_op('−')
            i += 1
            continue
        if c in '=<>':
            push_op(c)
            i += 1
            continue
        if c == '+':
            push_op('+')
            i += 1
            continue
        if c == '{':
            g, i = _read_group(s, i)
            push_runs(latex_to_runs(g, italic))
            continue
        if c == '}':
            i += 1
            continue
        if c == '\\':
            i += 1
            continue
        # 普通字符
        m = re.match(r'[A-Za-z]+', s[i:])
        if m:
            push(m.group(0), italic=italic)
            i += len(m.group(0))
            continue
        m = re.match(r'[0-9]+', s[i:])
        if m:
            push(m.group(0), italic=False)
            i += len(m.group(0))
            continue
        push(c, italic=False)
        i += 1
    return out


def _needs_paren(latex):
    t = re.sub(r'\\(text|mathrm|mathbf|mathit|operatorname)\{([^{}]*)\}', r'\2', latex)
    t = re.sub(r'\\sum|\\prod|\\int', 'S', t)
    return bool(re.search(r'(\s|\\,|\\;|\\times|\\cdot|\+|-|\\bar)', t))


# ---------------------------------------------------------------- HTML 片段解析

TOKEN_RE = re.compile(
    r'<span(?P<sattrs>[^>]*)>|(?P<spanend></span>)|(?P<bstart><strong>)|(?P<bend></strong>)'
    r'|(?P<sstart><sub>)|(?P<send></sub>)|(?P<br><br\s*/?>)|(?P<math>\\\(.+?\\\))',
    re.S)

STYLE_RE = re.compile(r'([a-zA-Z-]+)\s*:\s*([^;]+)')


def parse_style_attr(attrs):
    m = re.search(r'style\s*=\s*"([^"]*)"', attrs or '')
    if not m:
        m = re.search(r"style\s*=\s*'([^']*)'", attrs or '')
    if not m:
        return {}
    return dict((k.strip().lower(), v.strip()) for k, v in STYLE_RE.findall(m.group(1)))


def parse_inline(inner, base, theme):
    """返回 run 列表。base 为 {color, fontFamily, bold, italic}"""
    runs = []
    pos = 0
    state = {'color': base.get('color'), 'family': base.get('fontFamily')}
    stack = []
    bold = base.get('bold', False)
    italic_base = base.get('italic', False)
    sub = False

    def add(text):
        if text == '':
            return
        runs.append({'text': html.unescape(text), 'color': state['color'],
                     'fontFamily': state['family'], 'bold': bold, 'sub': sub,
                     'italic': italic_base})

    for m in TOKEN_RE.finditer(inner):
        add(inner[pos:m.start()])
        pos = m.end()
        tok = m.group(0)
        if m.group('math') is not None:
            mruns = latex_to_runs(m.group('math')[2:-2], True)
            for r in mruns:
                d = dict(r)
                d.setdefault('color', state['color'])
                d['fontFamily'] = {'latin': 'Times New Roman', 'ea': None}
                d['bold'] = bold or d.get('bold', False)
                d['math'] = True
                if sub:
                    d['sub'] = True
                runs.append(d)
        elif m.group('sattrs') is not None:
            stack.append(dict(state))
            st = parse_style_attr(m.group('sattrs'))
            if 'color' in st:
                state = dict(state)
                state['color'] = theme.color(st['color'])
            if 'font-family' in st:
                state = dict(state)
                state['family'] = st['font-family']
        elif m.group('spanend') is not None:
            if stack:
                state = stack.pop()
        elif m.group('bstart') is not None:
            bold = True
        elif m.group('bend') is not None:
            bold = base.get('bold', False)
        elif m.group('sstart') is not None:
            sub = True
        elif m.group('send') is not None:
            sub = False
        elif m.group('br') is not None:
            runs.append({'br': True})
    add(inner[pos:])
    return runs


PARA_RE = re.compile(r'<p([^>]*)>(.*?)</p>', re.S)


def parse_paragraphs(raw, theme):
    """把 DSL 的 text 字段解析为段落列表"""
    raw = raw.strip('\n')
    paras = PARA_RE.findall(raw)
    if not paras:
        paras = [('', raw)]
    out = []
    for attrs, inner in paras:
        st = parse_style_attr(attrs)
        align = st.get('text-align')
        sb = None
        if 'margin-top' in st:
            mm = re.match(r'(-?[\d.]+)', st['margin-top'])
            if mm:
                sb = float(mm.group(1))
        out.append({'attrs': st, 'align': align, 'space_before': sb, 'html': inner})
    return out


def has_visible(p):
    txt = re.sub(r'<[^>]+>', '', p['html'])
    txt = re.sub(r'\\\((.*?)\\\)', r'\1', txt, flags=re.S).strip()
    return txt != ''


# ---------------------------------------------------------------- pptx 构建

def add_text(slide, el, content, theme):
    x, y, w, h = el['bounds']
    style = theme.text_style(content['style']) if 'style' in content else {}
    merged = {}
    merged.update(style)
    for k in ('fontSize', 'color', 'bold', 'italic', 'lineHeight', 'letterSpacing',
              'fontFamily', 'align', 'wrap'):
        if k in content:
            merged[k] = content[k]
    if 'color' in merged:
        merged['color'] = theme.color(merged['color'])
    latin, ea = split_family(merged.get('fontFamily'))

    paras = parse_paragraphs(content['text'], theme)
    paras = [p for p in paras if has_visible(p) or True]
    if not paras:
        return None

    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = bool(merged.get('wrap', True))
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = ALIGN_V.get((merged.get('align') or ['left', 'top'])[1], MSO_ANCHOR.TOP)
    elem_align = ALIGN_H.get((merged.get('align') or ['left', 'top'])[0], PP_ALIGN.LEFT)

    base = {'color': merged.get('color'), 'fontFamily': merged.get('fontFamily'),
            'bold': merged.get('bold', False), 'italic': merged.get('italic', False)}

    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if p['align']:
            para.alignment = ALIGN_H.get(p['align'], elem_align)
        else:
            para.alignment = elem_align
        if merged.get('lineHeight'):
            para.line_spacing = merged['lineHeight']
        if p['space_before']:
            para.space_before = Pt(p['space_before'])
        runs = parse_inline(p['html'], base, theme)
        for r in runs:
            if r.get('br'):
                para.add_line_break()
                continue
            run = para.add_run()
            run.text = r['text']
            fs = merged.get('fontSize', 15)
            run.font.size = Pt(fs)
            if r.get('bold'):
                run.font.bold = True
            if r.get('italic'):
                run.font.italic = True
            col = theme.color(r.get('color')) if r.get('color') else merged.get('color')
            if col:
                run.font.color.rgb = RGBColor.from_string(col.lstrip('#').upper())
            fam = r.get('fontFamily')
            if isinstance(fam, dict) and fam.get('ea') is None:
                rl, re_ = 'Times New Roman', ea
            else:
                rl, re_ = split_family(fam if fam is not None else merged.get('fontFamily'))
            set_run_font(run, rl, re_)
            if merged.get('letterSpacing'):
                set_run_spacing(run, merged['letterSpacing'])
            if r.get('sub'):
                set_run_baseline(run, -25000)
            if r.get('sup'):
                set_run_baseline(run, 30000)
    return tb


def add_rect(slide, el, theme):
    x, y, w, h = el['bounds']
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    shp.shadow.inherit = False
    fill = el.get('fill')
    if fill and fill.get('type') == 'solid':
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor.from_string(theme.color(fill['color']).lstrip('#').upper())
    else:
        shp.fill.background()
    border = el.get('border')
    if border:
        shp.line.color.rgb = RGBColor.from_string(theme.color(border['color']).lstrip('#').upper())
        shp.line.width = Pt(border.get('width', 1))
    else:
        no_line(shp)
    shp.text_frame.text = ''
    return shp


def add_line(slide, el, theme):
    x, y, w, h = el['bounds']
    pts = [tuple(float(v) for v in p.split(',')) for p in el['points'].split()]
    (x1, y1), (x2, y2) = pts[0], pts[-1]
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x + x1), E(y + y1), E(x + x2), E(y + y2))
    border = el.get('border', {})
    conn.line.color.rgb = RGBColor.from_string(theme.color(border.get('color', '#000000')).lstrip('#').upper())
    conn.line.width = Pt(border.get('width', 1))
    arrows = el.get('arrow') or [None, None]
    ln = conn.line._get_or_add_ln()
    if len(arrows) > 0 and arrows[0] == 'arrow':
        he = OxmlElement('a:headEnd')
        he.set('type', 'triangle')
        he.set('w', 'med')
        he.set('len', 'med')
        ln.append(he)
    if len(arrows) > 1 and arrows[1] == 'arrow':
        te = OxmlElement('a:tailEnd')
        te.set('type', 'triangle')
        te.set('w', 'med')
        te.set('len', 'med')
        ln.append(te)
    return conn


def add_image(slide, el, base_dir, theme):
    from PIL import Image
    x, y, w, h = el['bounds']
    path = os.path.join(base_dir, el['src'].replace('/', os.sep))
    iw, ih = Image.open(path).size
    ar_img = float(iw) / float(ih)
    ar_box = float(w) / float(h)
    fit = (el.get('fit') or {}).get('mode', 'contain')
    if fit == 'contain':
        if ar_img > ar_box:
            nw, nh = w, w / ar_img
        else:
            nh, nw = h, h * ar_img
        pic = slide.shapes.add_picture(path, E(x + (w - nw) / 2.0), E(y + (h - nh) / 2.0), E(nw), E(nh))
    else:
        pic = slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))
        if ar_img > ar_box:
            c = (1 - ar_box / ar_img) / 2.0
            pic.crop_left = c
            pic.crop_right = c
        else:
            c = (1 - ar_img / ar_box) / 2.0
            pic.crop_top = c
            pic.crop_bottom = c
    border = el.get('border')
    if border:
        pic.line.color.rgb = RGBColor.from_string(theme.color(border['color']).lstrip('#').upper())
        pic.line.width = Pt(border.get('width', 1))
    else:
        no_line(pic)
    return pic


# ---------------------------------------------------------------- 表格

FILL_TAGS = ['a:noFill', 'a:solidFill', 'a:gradFill', 'a:blipFill', 'a:pattFill', 'a:grpFill']
BORDER_ORDER = ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB', 'a:lnTlToBr', 'a:lnBlToTr', 'a:cell3D']


def _tcPr(cell):
    return cell._tc.get_or_add_tcPr()


def cell_border(cell, edge, width_pt, color_hex):
    tcPr = _tcPr(cell)
    tag = 'a:' + edge
    for old in tcPr.findall(qn(tag)):
        tcPr.remove(old)
    ln = OxmlElement(tag)
    ln.set('w', str(int(round(width_pt * 12700))))
    ln.set('cap', 'flat')
    ln.set('cmpd', 'sng')
    ln.set('algn', 'ctr')
    if color_hex is None:
        ln.append(OxmlElement('a:noFill'))
    else:
        sf = OxmlElement('a:solidFill')
        clr = OxmlElement('a:srgbClr')
        clr.set('val', color_hex.lstrip('#').upper())
        sf.append(clr)
        ln.append(sf)
        pd = OxmlElement('a:prstDash')
        pd.set('val', 'solid')
        ln.append(pd)
    _insert_ordered(tcPr, ln, BORDER_ORDER)


def clear_cell_borders(cell):
    for e in ('lnL', 'lnR', 'lnT', 'lnB', 'lnTlToBr', 'lnBlToTr'):
        cell_border(cell, e, 0, None)


def set_table_style_none(table):
    tblPr = table._tbl.tblPr
    tblPr.set('firstRow', '0')
    tblPr.set('bandRow', '0')
    for el in tblPr.findall(qn('a:tableStyleId')):
        tblPr.remove(el)
    el = OxmlElement('a:tableStyleId')
    el.text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
    tblPr.append(el)


def add_table(slide, el, theme):
    x, y, w, h = el['bounds']
    rows = el['rows']
    nrow = len(rows)
    ncol = max(len(r) for r in rows)
    colfrac = el.get('columnWidths') or [1.0 / ncol] * ncol
    rowfrac = el.get('rowHeights') or [1.0 / nrow] * nrow
    gf = slide.shapes.add_table(nrow, ncol, E(x), E(y), E(w), E(h))
    table = gf.table
    set_table_style_none(table)
    for i, f in enumerate(colfrac):
        table.columns[i].width = E(w * f)
    for i, f in enumerate(rowfrac):
        table.rows[i].height = E(h * f)

    # 建立网格（处理 rowSpan）
    grid = [[None] * ncol for _ in range(nrow)]
    occupied = set()
    for r in range(nrow):
        cells = rows[r]
        idx = 0
        c = 0
        while idx < len(cells) and c < ncol:
            while c < ncol and (r, c) in occupied:
                c += 1
            if c >= ncol:
                break
            cell = cells[idx]
            idx += 1
            span = int(cell.get('rowSpan', 1) or 1)
            grid[r][c] = cell
            for rr in range(r + 1, min(r + span, nrow)):
                occupied.add((rr, c))
            c += 1

    # 合并
    for r in range(nrow):
        for c in range(ncol):
            cell = grid[r][c]
            if cell and int(cell.get('rowSpan', 1) or 1) > 1:
                span = int(cell['rowSpan'])
                table.cell(r, c).merge(table.cell(min(r + span, nrow) - 1, c))

    # 填充与样式
    base_fs = el.get('styleFontSize', 13)
    for r in range(nrow):
        for c in range(ncol):
            tc = table.cell(r, c)
            tc.fill.background()
            clear_cell_borders(tc)
            tc.margin_left = Pt(3)
            tc.margin_right = Pt(3)
            tc.margin_top = Pt(1)
            tc.margin_bottom = Pt(1)
            tc.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell = grid[r][c]
            if cell is None or cell.get('skip'):
                continue
            # 文本
            html_text = cell.get('text', '')
            fs = float(cell.get('fontSize', base_fs))
            latin, ea = split_family(cell.get('fontFamily'))
            bold = bool(cell.get('bold', False))
            color = theme.color(cell.get('color')) if cell.get('color') else None
            if r == 0:
                color = color or theme.color('$primary')
                bold = True
            align_attr = cell.get('align')
            if align_attr:
                halign = ALIGN_H.get(align_attr[0], PP_ALIGN.CENTER)
                tc.vertical_anchor = ALIGN_V.get(align_attr[1], MSO_ANCHOR.MIDDLE)
            else:
                halign = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
            if not (r == 0 and c == 0):
                pass
            tf = tc.text_frame
            tf.word_wrap = True
            paras = parse_paragraphs(html_text, theme)
            base = {'color': color or theme.color('$text'), 'fontFamily': cell.get('fontFamily'),
                    'bold': bold, 'italic': False}
            first = True
            for p in paras:
                if len(paras) > 1 and not has_visible(p):
                    continue
                para = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                para.alignment = ALIGN_H.get(p['align'], halign) if p['align'] else halign
                runs = parse_inline(p['html'], base, theme)
                for rr in runs:
                    if rr.get('br'):
                        para.add_line_break()
                        continue
                    run = para.add_run()
                    run.text = rr['text']
                    run.font.size = Pt(fs)
                    if rr.get('bold'):
                        run.font.bold = True
                    if rr.get('italic'):
                        run.font.italic = True
                    col = theme.color(rr.get('color')) if rr.get('color') else (color or theme.color('$text'))
                    if col:
                        run.font.color.rgb = RGBColor.from_string(col.lstrip('#').upper())
                    fam = rr.get('fontFamily')
                    if isinstance(fam, dict) and fam.get('ea') is None:
                        rl, re_ = 'Times New Roman', ea
                    else:
                        rl, re_ = split_family(fam if fam is not None else cell.get('fontFamily'))
                    set_run_font(run, rl, re_)
                    if rr.get('sub'):
                        set_run_baseline(run, -25000)
                    if rr.get('sup'):
                        set_run_baseline(run, 30000)

    # 三线表框线
    tcol = theme.color('$text')
    for c in range(ncol):
        cell_border(table.cell(0, c), 'lnT', 1.5, tcol)
        cell_border(table.cell(0, c), 'lnB', 1.0, tcol)
        cell_border(table.cell(nrow - 1, c), 'lnB', 1.5, tcol)
    return table


def set_slide_bg(slide, hexcolor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hexcolor.lstrip('#').upper())


# ---------------------------------------------------------------- 主流程

def build(src_dir, out_path):
    pptd_path = None
    for f in os.listdir(src_dir):
        if f.lower().endswith('.pptd'):
            pptd_path = os.path.join(src_dir, f)
            break
    with open(pptd_path, 'r', encoding='utf-8') as fh:
        pptd = yaml.safe_load(fh)
    theme = Theme(pptd)
    W, H = pptd.get('size', [960, 540])

    prs = Presentation()
    prs.slide_width = E(W)
    prs.slide_height = E(H)
    blank = prs.slide_layouts[6]

    stats = []
    for rel in pptd['pages']:
        page_path = os.path.join(src_dir, rel.replace('/', os.sep))
        with open(page_path, 'r', encoding='utf-8') as fh:
            page = yaml.safe_load(fh)
        slide = prs.slides.add_slide(blank)
        bg = (page.get('background') or {})
        set_slide_bg(slide, bg.get('color', '#FFFFFF') if bg.get('type') == 'solid' else '#FFFFFF')
        n_text = n_shape = n_img = n_tbl = n_line = 0
        for el in page.get('elements') or []:
            et = el['elementType']
            try:
                if et == 'text':
                    # 整体为一个公式的文本元素 -> 用公式图片渲染
                    c = el.get('content') or {}
                    txt = (c.get('text') or '').strip()
                    whole = re.fullmatch(r'<p[^>]*>\s*\\\((.*?)\\\)\s*</p>', txt, re.S)
                    if whole:
                        img = render_formula(whole.group(1), os.path.join(
                            os.environ.get('TEMP', '.'), '_formula_%s.png' % el['elementId']),
                            float(c.get('fontSize', 15)))
                        if img:
                            add_formula_image(slide, el, img, theme)
                            n_img += 1
                            continue
                    add_text(slide, el, c, theme)
                    n_text += 1
                elif et == 'shape':
                    add_rect(slide, el, theme)
                    n_shape += 1
                elif et == 'line':
                    add_line(slide, el, theme)
                    n_line += 1
                elif et == 'image':
                    add_image(slide, el, src_dir, theme)
                    n_img += 1
                elif et == 'table':
                    if 'styleFontSize' not in el:
                        el = dict(el)
                        st = theme.tableStyles.get((el.get('style') or '').lstrip('$'), {})
                        cs = st.get('cellStyle', {})
                        el['styleFontSize'] = cs.get('fontSize', 13)
                    add_table(slide, el, theme)
                    n_tbl += 1
                else:
                    raise ValueError('unknown elementType ' + et)
            except Exception as ex:
                import traceback
                traceback.print_exc()
                raise RuntimeError('page %s element %s: %r' % (rel, el.get('elementId'), ex))
        stats.append((rel, n_text, n_shape, n_line, n_img, n_tbl))

    prs.save(out_path)
    return stats


def add_formula_image(slide, el, png_path, theme):
    from PIL import Image
    x, y, w, h = el['bounds']
    with Image.open(png_path) as im:
        iw, ih = im.size
    # png 以 dpi=600 渲染：1pt = 600/72 px
    scale = 72.0 / 600.0
    lw, lh = iw * scale, ih * scale
    fit = min(1.0, w / lw, h / lh)
    nw, nh = lw * fit, lh * fit
    slide.shapes.add_picture(png_path, E(x + (w - nw) / 2.0), E(y + (h - nh) / 2.0), E(nw), E(nh))


def render_formula(latex, out_png, fontsize_pt):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib import font_manager
    try:
        font_manager.fontManager.addfont(r'C:\Windows\Fonts\simsun.ttc')
        font_manager.fontManager.addfont(r'C:\Windows\Fonts\times.ttf')
    except Exception:
        pass
    plt.rcParams['mathtext.fontset'] = 'custom'
    plt.rcParams['mathtext.rm'] = 'SimSun'
    plt.rcParams['mathtext.it'] = 'Times New Roman:italic'
    plt.rcParams['mathtext.bf'] = 'SimSun:bold'
    tex = re.sub(r'\\text\{([^{}]*)\}', lambda m: '\\mathrm{%s}' % m.group(1), latex)
    fig = plt.figure(figsize=(20, 4), dpi=600)
    fig.text(0.5, 0.5, '$%s$' % tex, ha='center', va='center', fontsize=fontsize_pt)
    fig.savefig(out_png, transparent=True, bbox_inches='tight', pad_inches=0.02)
    plt.close(fig)
    return out_png


if __name__ == '__main__':
    src = sys.argv[1]
    dst = sys.argv[2]
    st = build(src, dst)
    lines = []
    for s in st:
        lines.append('%-28s text=%-3d shape=%-3d line=%-2d image=%-2d table=%d' % s)
    with open(os.path.join(os.environ.get('TEMP', '.'), '_build_stats.txt'), 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(lines))
    print('OK ->', dst)
